"""Convierte un reporte (título, encabezados, filas) en un archivo descargable."""
import csv
import io
from datetime import date, datetime
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone
from django.utils.text import slugify

TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y %H:%M")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, Decimal):
        return f"{valor:.2f}"
    return str(valor)


def _nombre_archivo(titulo, extension):
    return f"{slugify(titulo)}_{timezone.now():%Y%m%d}.{extension}"


def respuesta_csv(titulo, encabezados, filas):
    """CSV en UTF-8 con BOM para que Excel muestre bien las tildes."""
    salida = io.StringIO()
    escritor = csv.writer(salida)
    escritor.writerow(encabezados)
    for fila in filas:
        escritor.writerow([_texto(valor) for valor in fila])

    respuesta = HttpResponse("﻿" + salida.getvalue(), content_type="text/csv; charset=utf-8")
    respuesta["Content-Disposition"] = f'attachment; filename="{_nombre_archivo(titulo, "csv")}"'
    return respuesta


def respuesta_xlsx(titulo, encabezados, filas, subtitulo=""):
    """Libro de Excel con título, encabezados con formato y columnas ajustadas."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo[:31]

    hoja.append([f"Óptica Moda Visual — {titulo}"])
    hoja["A1"].font = Font(bold=True, size=14, color="174785")
    hoja.append([subtitulo or f"Generado el {timezone.now():%d/%m/%Y %H:%M}"])
    hoja["A2"].font = Font(italic=True, color="718096")
    hoja.append([])
    hoja.append(encabezados)

    for celda in hoja[4]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="174785")
        celda.alignment = Alignment(horizontal="center")

    for fila in filas:
        hoja.append([float(v) if isinstance(v, Decimal) else v for v in fila])

    for indice, encabezado in enumerate(encabezados, start=1):
        letra = get_column_letter(indice)
        largo = max([len(encabezado)] + [len(_texto(f[indice - 1])) for f in filas])
        hoja.column_dimensions[letra].width = min(max(largo + 2, 10), 50)
        for celda in hoja[letra][4:]:
            if "(Q)" in encabezado:
                celda.number_format = '"Q" #,##0.00'
            elif isinstance(celda.value, datetime):
                celda.number_format = "dd/mm/yyyy hh:mm"
            elif isinstance(celda.value, date):
                celda.number_format = "dd/mm/yyyy"

    hoja.freeze_panes = "A5"

    salida = io.BytesIO()
    libro.save(salida)
    respuesta = HttpResponse(salida.getvalue(), content_type=TIPO_XLSX)
    respuesta["Content-Disposition"] = f'attachment; filename="{_nombre_archivo(titulo, "xlsx")}"'
    return respuesta
